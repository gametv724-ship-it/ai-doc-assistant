from sentence_transformers import SentenceTransformer
from qdrant_client import QdrantClient
from qdrant_client.models import Distance, VectorParams, PointStruct
import PyPDF2
from docx import Document
import openpyxl
import uuid
import os

class DocumentRAG:
    def __init__(self, use_docker=True):
        print("🔄 Загружаю embeddings модель...")
        self.embedder = SentenceTransformer('all-MiniLM-L6-v2')
        
        # Подключаемся к Qdrant
        self.use_docker = False  # По умолчанию память
        
        if use_docker:
            try:
                # Пробуем подключиться к Docker
                test_client = QdrantClient(host="localhost", port=6333, timeout=2)
                test_client.get_collections()  # Проверяем доступность
                self.qdrant = test_client
                self.use_docker = True
                print("✅ Подключился к Qdrant (Docker)")
            except Exception as e:
                print(f"⚠️ Docker недоступен ({str(e)[:50]}), использую память")
                self.qdrant = QdrantClient(":memory:")
        else:
            self.qdrant = QdrantClient(":memory:")
            print("✅ Использую Qdrant в памяти")
        
        self.collection_name = "documents"
        self._create_collection()
    
    def _create_collection(self):
        """Создаём коллекцию, если её нет"""
        try:
            collections = self.qdrant.get_collections().collections
            collection_exists = any(c.name == self.collection_name for c in collections)
            
            if not collection_exists:
                self.qdrant.create_collection(
                    collection_name=self.collection_name,
                    vectors_config=VectorParams(size=384, distance=Distance.COSINE)
                )
                print(f"✅ Создана коллекция '{self.collection_name}'")
            else:
                print(f"✅ Коллекция '{self.collection_name}' уже существует")
        except Exception as e:
            # Если коллекции нет - создаём
            self.qdrant.create_collection(
                collection_name=self.collection_name,
                vectors_config=VectorParams(size=384, distance=Distance.COSINE)
            )
            print(f"✅ Создана коллекция '{self.collection_name}'")
    
    def load_pdf(self, file_path):
        """Читаем PDF"""
        text = ""
        with open(file_path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
        return text
    
    def load_docx(self, file_path):
        """Читаем Word"""
        doc = Document(file_path)
        return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    
    def load_xlsx(self, file_path):
        """Читаем Excel"""
        wb = openpyxl.load_workbook(file_path)
        text = ""
        for sheet in wb.worksheets:
            text += f"\n=== {sheet.title} ===\n"
            for row in sheet.iter_rows(values_only=True):
                row_text = " | ".join([str(cell) if cell else "" for cell in row])
                if row_text.strip():
                    text += row_text + "\n"
        return text
    
    def chunk_text(self, text, chunk_size=500, overlap=50):
        """Режем на чанки"""
        words = text.split()
        chunks = []
        for i in range(0, len(words), chunk_size - overlap):
            chunk = " ".join(words[i:i + chunk_size])
            if chunk.strip():
                chunks.append(chunk)
        return chunks
    
    def add_document(self, file_path):
        """Загружаем документ в векторную БД"""
        ext = os.path.splitext(file_path)[1].lower()
        
        try:
            if ext == '.pdf':
                text = self.load_pdf(file_path)
            elif ext == '.docx':
                text = self.load_docx(file_path)
            elif ext in ['.xlsx', '.xls']:
                text = self.load_xlsx(file_path)
            else:
                return f"❌ Неподдерживаемый формат: {ext}"
            
            if not text.strip():
                return "⚠️ Документ пустой или не удалось извлечь текст"
            
            chunks = self.chunk_text(text)
            
            if not chunks:
                return "⚠️ Не удалось создать чанки из текста"
            
            print(f"📝 Создаю embeddings для {len(chunks)} чанков...")
            embeddings = self.embedder.encode(chunks, show_progress_bar=True)
            
            points = [
                PointStruct(
                    id=str(uuid.uuid4()),
                    vector=embedding.tolist(),
                    payload={
                        "text": chunk, 
                        "source": os.path.basename(file_path)
                    }
                )
                for chunk, embedding in zip(chunks, embeddings)
            ]
            
            self.qdrant.upsert(
                collection_name=self.collection_name,
                points=points
            )
            
            return f"✅ Загружено {len(chunks)} чанков из {os.path.basename(file_path)}"
            
        except Exception as e:
            return f"❌ Ошибка обработки файла: {str(e)}"
    
    def search(self, query, top_k=3):
        """Ищем похожие чанки"""
        try:
            query_vector = self.embedder.encode(query).tolist()
            
            results = self.qdrant.search(
                collection_name=self.collection_name,
                query_vector=query_vector,
                limit=top_k
            )
            
            return [
                {
                    "text": hit.payload["text"],
                    "score": hit.score,
                    "source": hit.payload["source"]
                }
                for hit in results
            ]
        except Exception as e:
            print(f"❌ Ошибка поиска: {e}")
            return []
    
    def answer_question(self, question, use_llm=False, groq_api_key=None):
        """Получаем ответ на вопрос"""
        context_chunks = self.search(question, top_k=3)
        
        if not context_chunks:
            return {
                "answer": "❌ Не нашёл информации в документе. Попробуй загрузить файл сначала.",
                "sources": []
            }
        
        context = "\n\n".join([chunk["text"] for chunk in context_chunks])
        
        # Если есть Groq API — используем LLM
        if use_llm and groq_api_key and groq_api_key.strip():
            try:
                from groq import Groq
                
                # Очищаем ключ от лишних пробелов
                api_key_clean = groq_api_key.strip()
                
                client = Groq(api_key=api_key_clean)
                
                response = client.chat.completions.create(
                    model="llama-3.1-8b-instant",
                    messages=[
                        {
                            "role": "system", 
                            "content": "Ты помощник по документам. Отвечай кратко и по делу на русском языке, используя только информацию из предоставленного контекста."
                        },
                        {
                            "role": "user", 
                            "content": f"Контекст из документа:\n{context}\n\nВопрос: {question}\n\nОтвет:"
                        }
                    ],
                    temperature=0.3,
                    max_tokens=500
                )
                
                answer = response.choices[0].message.content
                
            except Exception as e:
                # Возвращаем просто текст без упоминания ошибки
                answer = f"📄 Найденная информация:\n\n{context}"
                print(f"⚠️ Ошибка Groq API: {str(e)}")
        else:
            # Без LLM — просто вернём найденный текст
            answer = f"📄 Найденная информация:\n\n{context}"
        
        return {
            "answer": answer,
            "sources": context_chunks
        }
